
#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
数学题目生成器
"""

import random
import os
import json

script_dir = os.path.dirname(os.path.abspath(__file__))
skill_root = os.path.dirname(script_dir)
tmp_files_dir = os.path.join(skill_root, 'tmp_files')
if not os.path.exists(tmp_files_dir):
    os.makedirs(tmp_files_dir)


def _build_carry_addition_bank():
    bank = []
    for a in range(1, 20):
        for b in range(1, 20):
            total = a + b
            if 1 <= total <= 19:
                a_units = a % 10
                b_units = b % 10
                if a_units + b_units >= 10:
                    question = "%d + %d =" % (a, b)
                    bank.append((question, total))
    return bank


def _build_borrow_subtraction_bank():
    bank = []
    for a in range(10, 20):
        for b in range(1, 10):
            total = a - b
            if 1 <= total <= 18:
                a_units = a % 10
                if a_units < b:
                    question = "%d - %d =" % (a, b)
                    bank.append((question, total))
    return bank


def _sample_with_probability(bank, check_func, prob=0.3):
    while True:
        q, a = random.choice(bank)
        if check_func(q, a):
            if random.random() < prob:
                return q, a
        else:
            return q, a


def generate_carry_addition():
    def check_func(q, a):
        return a in (10, 11)

    return _sample_with_probability(_carry_add_bank, check_func, 0.3)


def generate_borrow_subtraction():
    def check_func(q, a):
        parts = q.split("-")
        b_num = int(parts[1].strip().split("=")[0].strip())
        return b_num in (1, 2, 3)

    return _sample_with_probability(_borrow_sub_bank, check_func, 0.3)


def _sample_from_bank(bank, count):
    bank_len = len(bank)
    if count <= bank_len:
        return random.sample(bank, count)
    else:
        result = bank.copy()
        remaining = count - bank_len
        for _ in range(remaining):
            result.append(random.choice(bank))
        return result


def generate_20_mixed_add_sub(count=100):
    half_count = count // 2
    add_questions = _sample_from_bank(_carry_add_bank, half_count)
    sub_questions = _sample_from_bank(_borrow_sub_bank, count - half_count)

    questions = []
    for q in add_questions:
        questions.append(q)
    for q in sub_questions:
        questions.append(q)

    random.shuffle(questions)
    return questions


def format_questions_grid(questions, per_row=4, show_answers=False):
    lines = []
    for i in range(0, len(questions), per_row):
        row_questions = questions[i:i + per_row]
        row_str = ""
        for q, a in row_questions:
            if show_answers:
                row_str += "%s %-6s" % (q, str(a))
            else:
                row_str += "%-10s" % q
        lines.append(row_str)
    return "\n".join(lines)


def save_to_docx(questions, output_file, per_row=4, show_answers=False):
    try:
        from docx import Document
        from docx.shared import Pt, Inches

        doc = Document()

        title = doc.add_heading("20以内进退位加减法练习", 0)
        title.alignment = 1

        p = doc.add_paragraph()
        p.add_run("姓名：__________  班级：__________  日期：__________  用时：__________")
        p.alignment = 1
        p.add_run("\n\n")

        table = doc.add_table(rows=(len(questions) + per_row - 1) // per_row, cols=per_row)
        table.style = "Table Grid"

        table_width = Inches(6.5)
        column_width = table_width / per_row
        for column in table.columns:
            for cell in column.cells:
                cell.width = column_width

        for row_idx in range((len(questions) + per_row - 1) // per_row):
            for col_idx in range(per_row):
                q_idx = row_idx * per_row + col_idx
                cell = table.rows[row_idx].cells[col_idx]

                if q_idx < len(questions):
                    q, a = questions[q_idx]
                    if show_answers:
                        cell.text = "%s %s" % (q, a)
                    else:
                        cell.text = q
                else:
                    cell.text = ""

                paragraph = cell.paragraphs[0]
                paragraph.alignment = 0
                for run in paragraph.runs:
                    run.font.size = Pt(16)

        doc.save(output_file)
        return True
    except ImportError:
        print("提示: python-docx 库未安装，将保存为 txt 格式")
        txt_content = format_questions_grid(questions, per_row, show_answers)
        with open(output_file.replace(".docx", ".txt"), "w", encoding="utf-8") as f:
            f.write(txt_content)
        return False


def get_range_by_difficulty(difficulty, op_type):
    if op_type in ["add", "sub"]:
        ranges = {
            1: (1, 10),
            2: (1, 20),
            3: (1, 100),
            4: (10, 500),
            5: (100, 1000)
        }
    elif op_type == "mul":
        ranges = {
            1: (1, 5),
            2: (1, 9),
            3: (1, 9),
            4: (2, 12),
            5: (10, 99)
        }
    elif op_type == "div":
        ranges = {
            1: (2, 5),
            2: (2, 9),
            3: (2, 9),
            4: (2, 12),
            5: (10, 99)
        }
    else:
        ranges = {
            1: (1, 10),
            2: (1, 20),
            3: (1, 100),
            4: (1, 500),
            5: (1, 1000)
        }
    return ranges.get(difficulty, (1, 20))


def generate_addition(difficulty=2):
    min_val, max_val = get_range_by_difficulty(difficulty, "add")
    a = random.randint(min_val, max_val)
    b = random.randint(min_val, max_val)
    return "%d + %d =" % (a, b), a + b


def generate_subtraction(difficulty=2):
    min_val, max_val = get_range_by_difficulty(difficulty, "sub")
    a = random.randint(min_val, max_val)
    b = random.randint(min_val, max_val)
    if a < b:
        a, b = b, a
    return "%d - %d =" % (a, b), a - b


def generate_multiplication(difficulty=3):
    min_val, max_val = get_range_by_difficulty(difficulty, "mul")
    a = random.randint(min_val, max_val)
    b = random.randint(min_val, max_val)
    return "%d * %d =" % (a, b), a * b


def generate_division(difficulty=3):
    min_val, max_val = get_range_by_difficulty(difficulty, "div")
    b = random.randint(min_val, max_val)
    answer = random.randint(min_val, max_val)
    a = b * answer
    return "%d / %d =" % (a, b), answer


def generate_mixed(difficulty=3):
    ops = [
        ("+", generate_addition),
        ("-", generate_subtraction),
        ("*", generate_multiplication),
        ("/", generate_division)
    ]
    op_symbol, generator = random.choice(ops)
    return generator(difficulty)


def generate_questions(question_type="add", count=10, difficulty=2):
    generators = {
        "add": generate_addition,
        "sub": generate_subtraction,
        "mul": generate_multiplication,
        "div": generate_division,
        "mixed": generate_mixed
    }

    generator_func = generators.get(question_type, generate_addition)
    questions = []

    for _ in range(count):
        questions.append(generator_func(difficulty))

    return questions


def format_questions(questions, show_answers=False):
    lines = []
    for i, (question, answer) in enumerate(questions, 1):
        if show_answers:
            lines.append("%d. %s %d" % (i, question, answer))
        else:
            lines.append("%d. %s" % (i, question))
    return "\n".join(lines)


def _save_bank_to_file(bank, filename):
    filepath = os.path.join(tmp_files_dir, filename)
    with open(filepath, 'w', encoding='utf-8') as f:
        json.dump(bank, f, indent=2, ensure_ascii=False)


def _load_bank_from_file(filename):
    filepath = os.path.join(tmp_files_dir, filename)
    if os.path.exists(filepath):
        with open(filepath, 'r', encoding='utf-8') as f:
            return json.load(f)
    return None


def _init_banks():
    add_bank = _load_bank_from_file('carry_add_bank.json')
    sub_bank = _load_bank_from_file('borrow_sub_bank.json')
    
    if add_bank is None:
        add_bank = _build_carry_addition_bank()
        _save_bank_to_file(add_bank, 'carry_add_bank.json')
    
    if sub_bank is None:
        sub_bank = _build_borrow_subtraction_bank()
        _save_bank_to_file(sub_bank, 'borrow_sub_bank.json')
    
    return add_bank, sub_bank


_carry_add_bank, _borrow_sub_bank = _init_banks()
