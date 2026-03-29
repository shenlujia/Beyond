
#!/usr/bin/env python3
import os
import sys
import random
from datetime import datetime

from generator import generate_20_mixed_add_sub, save_to_docx

script_dir = os.path.dirname(os.path.abspath(__file__))
skill_root = os.path.dirname(script_dir)
output_dir = os.path.join(skill_root, 'tmp_files')

if not os.path.exists(output_dir):
    os.makedirs(output_dir)
    print('Created folder:', output_dir)

timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
filename = os.path.join(output_dir, '20以内进退位加减法_3页_100题每页_%s.docx' % timestamp)

print('Generating document with timestamp:', timestamp)

try:
    from docx import Document
    from docx.shared import Pt, Inches
    
    doc = Document()
    
    for page_num in range(1, 4):
        questions = generate_20_mixed_add_sub(100)
        
        table = doc.add_table(rows=(len(questions) + 4 - 1) // 4, cols=4)
        table.style = 'Table Grid'
        
        table_width = Inches(6.5)
        column_width = table_width / 4
        for column in table.columns:
            for cell in column.cells:
                cell.width = column_width
        
        for row_idx in range((len(questions) + 4 - 1) // 4):
            for col_idx in range(4):
                q_idx = row_idx * 4 + col_idx
                cell = table.rows[row_idx].cells[col_idx]
                
                if q_idx < len(questions):
                    q, a = questions[q_idx]
                    cell.text = q
                else:
                    cell.text = ""
                
                paragraph = cell.paragraphs[0]
                paragraph.alignment = 0
                for run in paragraph.runs:
                    run.font.size = Pt(16)
        
        if page_num < 3:
            doc.add_page_break()
    
    doc.save(filename)
    print('Document saved to:', filename)
    
except ImportError:
    print('提示: python-docx 库未安装')
